
# -*- coding: utf-8 -*-
from docx import Document
import sys

# 重定向输出到UTF-8
sys.stdout.reconfigure(encoding='utf-8')

doc = Document(r'C:\Users\32891\Downloads\2026_4_28133057_b7f.docx')

# 输出段落
for i, p in enumerate(doc.paragraphs):
    print(p.text)

# 输出表格
print("\n=== TABLES ===")
for table in doc.tables:
    for row in table.rows:
        print("|".join([cell.text for cell in row.cells]))
