
# -*- coding: utf-8 -*-
from docx import Document

# 打开文档
doc = Document(r'C:\Users\32891\Desktop\终稿_李莹莹.docx')

# 保存提取的内容
with open(r'C:\Users\32891\.qclaw\workspace\终稿_李莹莹_提取.txt', 'w', encoding='utf-8') as f:
    # 写入段落
    f.write('='*80 + '\n')
    f.write('文档内容\n')
    f.write('='*80 + '\n\n')
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if text:
            f.write(f'{text}\n')
    
    # 写入表格
    f.write('\n\n' + '='*80 + '\n')
    f.write('表格内容\n')
    f.write('='*80 + '\n\n')
    for i, table in enumerate(doc.tables):
        f.write(f'\n--- 表格 {i+1}: {len(table.rows)}行{len(table.columns)}列 ---\n')
        for row in table.rows:
            row_text = '\t'.join([cell.text.strip() for cell in row.cells])
            f.write(row_text + '\n')

print('文档提取完成！已保存为：终稿_李莹莹_提取.txt')
