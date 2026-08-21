import openpyxl, xlrd, csv, statistics
from collections import defaultdict

# ======== 1. 读取模板 ========
wb2 = xlrd.open_workbook(r'C:\Users\32891\Desktop\成绩导入模板.xls')
ws2 = wb2.sheet_by_index(0)
template = []
for r in range(2, ws2.nrows):
    raw_sid = ws2.cell_value(r, 0)
    # xlrd 读数字型学号会返回 float，转为整数字符串
    if isinstance(raw_sid, float):
        sid = str(int(raw_sid))
    else:
        sid = str(raw_sid).strip()
    name = str(ws2.cell_value(r, 1)).strip()
    cls = str(ws2.cell_value(r, 2)).strip()
    if sid:
        template.append({'xl_row': r + 1, 'sid': sid, 'name': name, 'cls': cls})
print(f"模板学生: {len(template)}人")

# ======== 2. 读取AI通识百分制成绩 ========
wb3 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\25-26-2人工智能通识考试成绩.xlsx')
ws3 = wb3['学生信息']
ai_scores = {}
for row in ws3.iter_rows(min_row=2, values_only=True):
    sid = str(row[0]).strip()
    pct = row[6]
    status = row[4]
    if sid and pct and isinstance(pct, (int, float)):
        ai_scores[sid] = {'name': str(row[1]).strip(), 'pct': int(pct), 'status': str(status).strip()}
print(f"AI通识有效百分制: {len(ai_scores)}条")

# ======== 3. 读取实训成绩（按学生聚合） ========
wb1 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\导出实训成绩-20260802-816690146199515136.xlsx')
ws1 = wb1['sheet1']

student_cases = defaultdict(list)  # sid -> [{score, grade, case}]
for row in ws1.iter_rows(min_row=2, values_only=True):
    name, sid, case_name, practice, realtime, status, score, grade = row
    if status == '已确认':
        # score 可能是字符串如 '76.0' 或数字
        try:
            score_val = float(score)
            student_cases[sid].append({
                'name': name,
                'case': case_name,
                'score': score_val,
                'grade': grade
            })
        except (ValueError, TypeError):
            pass

print(f"DEBUG: 实训匹配后 student_cases 人数: {len(student_cases)}")

# 计算每个学生的平均分
train_avg = {}
for sid, cases in student_cases.items():
    scores = [c['score'] for c in cases]
    train_avg[sid] = {
        'name': cases[0]['name'],
        'avg': round(statistics.mean(scores), 1),
        'count': len(scores),
        'scores': scores,
        'min': min(scores),
        'max': max(scores)
    }

print(f"有实训成绩的学生: {len(train_avg)}人，共{sum(v['count'] for v in train_avg.values())}条已确认记录")

# ======== 4. 匹配并输出 ========
results = []
matched_ai = 0
matched_train = 0
both = 0
none_ai = []
none_train = []

for s in template:
    sid = s['sid']
    ai = ai_scores.get(sid)
    tr = train_avg.get(sid)

    ai_pct = ai['pct'] if ai else None
    ai_status = ai['status'] if ai else None
    train_avg_score = tr['avg'] if tr else None
    train_count = tr['count'] if tr else 0

    if ai:
        matched_ai += 1
    if tr:
        matched_train += 1
    if ai and tr:
        both += 1
    if not ai:
        none_ai.append(s)
    if not tr:
        none_train.append(s)

    # 建议录入分数: 优先AI通识成绩，无则用实训均分
    suggested = ai_pct if ai_pct is not None else (train_avg_score if train_avg_score is not None else '')

    results.append({
        'xl_row': s['xl_row'],
        'sid': sid,
        'name': s['name'],
        'cls': s['cls'],
        'ai_pct': ai_pct if ai_pct is not None else '',
        'ai_status': ai_status if ai_status else '',
        'train_avg': train_avg_score if train_avg_score is not None else '',
        'train_count': train_count if train_count > 0 else '',
        'suggested': round(suggested) if suggested != '' else '',
    })

print(f"\n匹配结果:")
print(f"  有AI通识成绩: {matched_ai}/{len(template)}")
print(f"  有实训成绩: {matched_train}/{len(template)}")
print(f"  两者都有: {both}")
print(f"  无AI通识成绩: {len(none_ai)} -> 建议标记缺考")
print(f"  无实训成绩: {len(none_train)}")

if none_ai:
    print("\n无AI通识成绩学生（建议标记缺考或待补充）:")
    for s in none_ai:
        print(f"  {s['sid']} {s['name']} ({s['cls']})")

# ======== 5. 统计 ========
ai_list = [r['ai_pct'] for r in results if r['ai_pct'] != '']
train_list = [r['train_avg'] for r in results if r['train_avg'] != '']
print(f"\n=== 成绩统计 ===")
if ai_list:
    print(f"AI通识成绩: n={len(ai_list)}, min={min(ai_list)}, max={max(ai_list)}, avg={statistics.mean(ai_list):.1f}")
if train_list:
    print(f"实训均分:   n={len(train_list)}, min={min(train_list):.1f}, max={max(train_list):.1f}, avg={statistics.mean(train_list):.1f}")

# ======== 6. 输出CSV ========
csv_path = r'C:\Users\32891\Desktop\实训成绩_百分制_合并结果.csv'
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=['xl行号', '学号', '姓名', '班级', 'AI通识百分制成绩', 'AI通识状态', '实训均分', '实训已确认案例数', '建议录入分数', '备注'])
    writer.writeheader()
    for r in results:
        suggested_val = r['suggested']
        note = ''
        if r['ai_pct'] == '' and r['train_avg'] == '':
            note = '两者均无数据'
        elif r['ai_pct'] == '':
            note = '无AI通识成绩，建议标记缺考'
        elif r['train_avg'] == '':
            note = '无实训成绩'
        writer.writerow({
            'xl行号': r['xl_row'],
            '学号': r['sid'],
            '姓名': r['name'],
            '班级': r['cls'],
            'AI通识百分制成绩': r['ai_pct'],
            'AI通识状态': r['ai_status'],
            '实训均分': r['train_avg'],
            '实训已确认案例数': r['train_count'],
            '建议录入分数': suggested_val,
            '备注': note,
        })

print(f"\nCSV已输出: {csv_path}")
