import openpyxl, xlrd, xlwt, csv, statistics
from collections import defaultdict

# ======== 1. 读取模板 ========
wb2 = xlrd.open_workbook(r'C:\Users\32891\Desktop\成绩导入模板.xls')
ws2 = wb2.sheet_by_index(0)
template = []
for r in range(2, ws2.nrows):
    raw_sid = ws2.cell_value(r, 0)
    sid = str(int(raw_sid)) if isinstance(raw_sid, float) else str(raw_sid).strip()
    name = str(ws2.cell_value(r, 1)).strip()
    cls = str(ws2.cell_value(r, 2)).strip()
    if sid:
        template.append({'xl_row': r + 1, 'sid': sid, 'name': name, 'cls': cls})
print(f"模板学生: {len(template)}人")

# ======== 2. 读取AI通识百分制 ========
wb3 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\25-26-2人工智能通识考试成绩.xlsx')
ws3 = wb3['学生信息']
ai_scores = {}
for row in ws3.iter_rows(min_row=2, values_only=True):
    sid = str(row[0]).strip()
    pct = row[6]
    if sid and pct and isinstance(pct, (int, float)):
        ai_scores[sid] = {'name': str(row[1]).strip(), 'pct': int(pct), 'status': str(row[4]).strip()}

# ======== 3. 读取实训 ========
wb1 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\导出实训成绩-20260802-816690146199515136.xlsx')
ws1 = wb1['sheet1']
student_cases = defaultdict(list)
for row in ws1.iter_rows(min_row=2, values_only=True):
    name, sid, case_name, practice, realtime, status, score, grade = row
    if status == '已确认':
        try:
            student_cases[sid].append({'name': name, 'score': float(score), 'grade': grade})
        except (ValueError, TypeError):
            pass

train_avg = {}
for sid, cases in student_cases.items():
    scores = [c['score'] for c in cases]
    train_avg[sid] = {'name': cases[0]['name'], 'avg': round(statistics.mean(scores), 1), 'count': len(scores)}

# ======== 4. 合并数据 ========
results = []
for s in template:
    sid = s['sid']
    ai = ai_scores.get(sid)
    tr = train_avg.get(sid)
    results.append({
        'xl_row': s['xl_row'], 'sid': sid, 'name': s['name'], 'cls': s['cls'],
        'ai': ai, 'tr': tr,
    })

# ======== 5. 用 xlwt 写入新的 .xls ========
wb_out = xlwt.Workbook(encoding='utf-8')
ws_out = wb_out.add_sheet('学生成绩信息统计表')

# 写表头说明行
ws_out.write(0, 0, xlrd.decode_formula(ws2.cell_value(0, 0)) if False else ws2.cell_value(0, 0))

# 写表头
for c_idx in range(ws2.ncols):
    ws_out.write(1, c_idx, ws2.cell_value(1, c_idx))

# 写学生数据
for s in results:
    xl_r = s['xl_row'] - 1  # xlrd 0-index
    ai = s['ai']
    tr = s['tr']

    # 列0: 学号, 列1: 姓名, 列2: 班级
    ws_out.write(xl_r, 0, s['sid'])
    ws_out.write(xl_r, 1, s['name'])
    ws_out.write(xl_r, 2, s['cls'])

    # 列3: 平时成绩 -> 留空（模板只有这些列）
    # 列4: 期末成绩
    # 列5: 实验成绩
    # 列6: 总成绩 -> 建议录入分数
    # 列7: 未通过原因

    if ai:
        ws_out.write(xl_r, 6, ai['pct'])
    elif tr:
        ws_out.write(xl_r, 6, tr['avg'])

    # 标记无AI成绩
    if not ai and not tr:
        ws_out.write(xl_r, 7, '待补充')
    elif not ai:
        ws_out.write(xl_r, 7, '缺AI通识成绩')

out_xls = r'C:\Users\32891\Desktop\成绩导入模板_已填分数.xls'
wb_out.save(out_xls)
print(f"\n模板已填充，输出: {out_xls}")

# ======== 6. 输出详细CSV ========
csv_path = r'C:\Users\32891\Desktop\实训成绩_百分制_合并结果.csv'
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=['学号', '姓名', '班级', 'AI通识百分制', 'AI状态', '实训均分', '实训案例数', '建议录入分数', '备注'])
    writer.writeheader()
    for s in results:
        ai = s['ai']
        tr = s['tr']
        ai_pct = ai['pct'] if ai else None
        tr_avg = tr['avg'] if tr else None
        suggested = ai_pct if ai_pct is not None else tr_avg
        note = ''
        if not ai_pct and not tr_avg:
            note = '两者均无数据'
        elif not ai_pct:
            note = '无AI通识成绩，建议标记缺考'
        elif not tr_avg:
            note = '无实训成绩'
        writer.writerow({
            '学号': s['sid'], '姓名': s['name'], '班级': s['cls'],
            'AI通识百分制': ai_pct if ai_pct is not None else '',
            'AI状态': ai['status'] if ai else '',
            '实训均分': tr_avg if tr_avg is not None else '',
            '实训案例数': tr['count'] if tr else '',
            '建议录入分数': round(suggested) if suggested is not None else '',
            '备注': note,
        })
print(f"详细CSV: {csv_path}")

# ======== 7. 统计汇总 ========
ai_list = [s['ai']['pct'] for s in results if s['ai']]
tr_list = [s['tr']['avg'] for s in results if s['tr']]
print(f"\n=== 汇总 ===")
print(f"AI通识成绩: n={len(ai_list)}, min={min(ai_list)}, max={max(ai_list)}, avg={statistics.mean(ai_list):.1f}")
print(f"实训均分:   n={len(tr_list)}, min={min(tr_list):.1f}, max={max(tr_list):.1f}, avg={statistics.mean(tr_list):.1f}")
print(f"仅有AI成绩: {sum(1 for s in results if s['ai'] and not s['tr'])}人")
print(f"仅有实训:   {sum(1 for s in results if not s['ai'] and s['tr'])}人")
print(f"两者都有:   {sum(1 for s in results if s['ai'] and s['tr'])}人")
print(f"两者都无:   {sum(1 for s in results if not s['ai'] and not s['tr'])}人")
