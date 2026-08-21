import sys
sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
doc = Document(r"C:\Users\32891\Desktop\6_转专业方案_temp.docx")

text_parts = []
for para in doc.paragraphs:
    text_parts.append(para.text)

for i, table in enumerate(doc.tables):
    text_parts.append(f"\n=== 表格 {i+1} ===")
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        text_parts.append(" | ".join(cells))

output = "\n".join(text_parts)

# Write to file with UTF-8 BOM for Windows compatibility
with open(r"C:\Users\32891\Desktop\6_转专业方案_正文.txt", "w", encoding="utf-8-sig") as f:
    f.write(output)

print(f"共 {len(text_parts)} 段落/表格行")
print("写入完成")
