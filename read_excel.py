import openpyxl, xlrd
from collections import defaultdict

# ========== 深入分析实训成绩文件 ==========
wb1 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\导出实训成绩-20260802-816690146199515136.xlsx')
ws1 = wb1['sheet1']

all_rows = []
for row in ws1.iter_rows(min_row=2, values_only=True):
    all_rows.append(row)

print(f"总行数: {len(all_rows)}")

# 统计有多少独立学生
students = {}
for r in all_rows:
    name, sid = r[0], r[1]
    if sid not in students:
        students[sid] = {'name': name, 'cases': 0, 'confirmed': 0}
    if r[5] == '已确认':
        students[sid]['confirmed'] += 1
    students[sid]['cases'] += 1

print(f"独立学生数: {len(students)}")

# 确认过的学生
confirmed_students = {sid: s for sid, s in students.items() if s['confirmed'] > 0}
print(f"有已确认成绩的学生数: {len(confirmed_students)}")

# 看几个确认过的学生
print("\n有成绩的学生样本（前5个）:")
for sid, s in list(confirmed_students.items())[:5]:
    print(f"  {s['name']} {sid}: {s['confirmed']}/{s['cases']} 个案例已确认")

# 看一下"已确认"行的案例成绩是否真的是数值
confirmed_scores = []
for r in all_rows:
    if r[5] == '已确认':
        score = r[6]
        grade = r[7]
        confirmed_scores.append((r[0], r[1], r[2], score, grade))

print(f"\n已确认行数: {len(confirmed_scores)}")
print("\n前10条已确认记录:")
for item in confirmed_scores[:10]:
    print(f"  {item}")

# ========== 模板中的未匹配学生 ==========
print("\n\n=== 未匹配学生详情 ===")
wb2 = xlrd.open_workbook(r'C:\Users\32891\Desktop\成绩导入模板.xls')
ws2 = wb2.sheet_by_index(0)

students_template = []
for r in range(2, ws2.nrows):
    sid = str(ws2.cell_value(r, 0)).strip()
    name = str(ws2.cell_value(r, 1)).strip()
    cls = str(ws2.cell_value(r, 2)).strip()
    if sid:
        students_template.append({'row': r, 'sid': sid, 'name': name, 'cls': cls})

# AI成绩
wb3 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\25-26-2人工智能通识考试成绩.xlsx')
ws3_info = wb3['学生信息']
ai_scores = {}
for row in ws3_info.iter_rows(min_row=2, values_only=True):
    sid = str(row[0]).strip()
    pct = row[6]
    if sid and pct and isinstance(pct, (int, float)):
        ai_scores[sid] = {'name': str(row[1]).strip(), 'pct': pct}

print("未匹配学生及姓名:")
for s in students_template:
    if s['sid'] not in ai_scores:
        print(f"  {s['sid']} {s['name']} ({s['cls']}) row={s['row']}")

# ========== 看AI通识文件中学号开头匹配情况 ==========
# 2025595011xx 在AI文件中搜索
prefix = '2025595011'
matching = {k: v for k, v in ai_scores.items() if k.startswith(prefix)}
print(f"\nAI文件中学号以 {prefix} 开头的人数: {len(matching)}")
for sid, v in list(matching.items())[:5]:
    print(f"  {sid} {v['name']} 成绩={v['pct']}")

prefix2 = '202559501'
matching2 = {k: v for k, v in ai_scores.items() if k.startswith(prefix2)}
print(f"\nAI文件中学号以 {prefix2} 开头的人数: {len(matching2)}")
