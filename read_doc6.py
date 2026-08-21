
# -*- coding: utf-8 -*-
import sys
from docx import Document

# 重定向输出到UTF-8
sys.stdout.reconfigure(encoding='utf-8')

try:
    # 先尝试用docx库直接读取
    doc = Document(r'C:\Users\32891\Downloads\2026_4_22185621_smt.doc')
    
    # 输出段落
    print("=== PARAGRAPHS ===")
    for i, p in enumerate(doc.paragraphs):
        print(p.text)
    
    # 输出表格
    print("\n=== TABLES ===")
    for table in doc.tables:
        for row in table.rows:
            print("|".join([cell.text for cell in row.cells]))
            
except Exception as e:
    print(f"Error: {e}")
    print("\n尝试使用win32com读取...")
    try:
        import win32com.client as win32
        word = win32.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(r"C:\Users\32891\Downloads\2026_4_22185621_smt.doc")
        print(doc.Content.Text)
        doc.Close()
        word.Quit()
    except Exception as e2:
        print(f"Error2: {e2}")
