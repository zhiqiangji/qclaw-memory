import openpyxl, xlrd
from collections import defaultdict

# 诊断：三个文件各自的学生班级分布
print("=" * 60)
print("1. 实训文件学生班级分布")
print("=" * 60)
wb1 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\导出实训成绩-20260802-816690146199515136.xlsx')
ws1 = wb1['sheet1']
students_train = {}
for row in ws1.iter_rows(min_row=2, values_only=True):
    name, sid = row[0], str(row[1])
    if sid not in students_train:
        students_train[sid] = name
print(f"独立学生: {len(students_train)} 人")
print("学号前缀分布:")
prefixes = defaultdict(int)
for sid in students_train:
    prefixes[sid[:8] + 'xxxx'] += 1
for p, n in sorted(prefixes.items()):
    print(f"  {p}: {n}人")
# 看几个样本
sample = list(students_train.items())[:5]
for sid, name in sample:
    print(f"  {name} {sid}")

print()
print("=" * 60)
print("2. 模板文件学生班级分布")
print("=" * 60)
wb2 = xlrd.open_workbook(r'C:\Users\32891\Desktop\成绩导入模板.xls')
ws2 = wb2.sheet_by_index(0)
students_tmpl = {}
for r in range(2, ws2.nrows):
    sid = str(ws2.cell_value(r, 0)).strip()
    name = str(ws2.cell_value(r, 1)).strip()
    cls = str(ws2.cell_value(r, 2)).strip()
    if sid:
        students_tmpl[sid] = {'name': name, 'cls': cls}
print(f"独立学生: {len(students_tmpl)} 人")
by_cls = defaultdict(list)
for sid, info in students_tmpl.items():
    by_cls[info['cls']].append(sid)
for cls, sids in sorted(by_cls.items()):
    print(f"  {cls}: {len(sids)}人 — 前3个: {sids[:3]}")

print()
print("=" * 60)
print("3. 模板 vs 实训：学号交集")
print("=" * 60)
train_sids = set(students_train.keys())
tmpl_sids = set(students_tmpl.keys())
overlap = train_sids & tmpl_sids
print(f"模板学号数: {len(tmpl_sids)}")
print(f"实训学号数: {len(train_sids)}")
print(f"交集: {len(overlap)} 个")
if overlap:
    for sid in list(overlap)[:5]:
        print(f"  共有: {students_tmpl[sid]['name']} {sid}")

print()
print("=" * 60)
print("4. AI通识文件：土251学生成绩")
print("=" * 60)
wb3 = openpyxl.load_workbook(r'C:\Users\32891\Desktop\25-26-2人工智能通识考试成绩.xlsx')
ws3 = wb3['学生信息']
ai土251 = {}
for row in ws3.iter_rows(min_row=2, values_only=True):
    sid = str(row[0]).strip()
    name = str(row[1]).strip()
    pct = row[6]
    status = row[4]
    if sid.startswith('2025595011') and pct and isinstance(pct, (int, float)):
        ai土251[sid] = {'name': name, 'pct': int(pct), 'status': str(status).strip()}
print(f"AI文件中土251学生（2025595011xx）: {len(ai土251)} 人")
by_range = defaultdict(list)
for sid, info in ai土251.items():
    by_range[sid[:10]].append(info)
for r, infos in sorted(by_range.items()):
    print(f"  {r}xx: {len(infos)}人")
    for i in infos[:3]:
        print(f"    {i['name']} 成绩={i['pct']} ({i['status']})")
