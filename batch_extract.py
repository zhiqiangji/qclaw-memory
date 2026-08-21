# -*- coding: utf-8 -*-
import sys, os, glob
sys.stdout.reconfigure(encoding='utf-8')

# Step 1: Convert .doc files to .docx via win32com
doc_files = glob.glob(r"C:\Users\32891\Desktop\审稿\2026_5_*.doc")
if doc_files:
    import win32com.client
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    for f in doc_files:
        out = f.replace(".doc", "_converted.docx")
        try:
            d = word.Documents.Open(f)
            d.SaveAs(out, FileFormat=16)
            d.Close()
            print(f"Converted: {os.path.basename(f)}")
        except Exception as e:
            print(f"Failed: {os.path.basename(f)} - {e}")
    try:
        word.Quit()
    except:
        pass

# Step 2: Extract text from all .docx files
from docx import Document

output_dir = r"C:\Users\32891\.qclaw\workspace\审稿_extracted"
os.makedirs(output_dir, exist_ok=True)

docx_files = glob.glob(r"C:\Users\32891\Desktop\审稿\2026_5_*.docx")
# Also include converted files
docx_files += glob.glob(r"C:\Users\32891\Desktop\审稿\*_converted.docx")

for f in docx_files:
    basename = os.path.splitext(os.path.basename(f))[0]
    try:
        doc = Document(f)
        text_parts = []
        for para in doc.paragraphs:
            text_parts.append(para.text)
        for i, table in enumerate(doc.tables):
            text_parts.append(f"\n=== 表格 {i+1} ===")
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                text_parts.append(" | ".join(cells))
        
        outpath = os.path.join(output_dir, basename + ".txt")
        with open(outpath, "w", encoding="utf-8-sig") as fout:
            fout.write("\n".join(text_parts))
        print(f"OK: {basename} ({len(text_parts)} lines)")
    except Exception as e:
        print(f"FAIL: {basename} - {e}")
