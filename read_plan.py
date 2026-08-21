import docx
doc = docx.Document(r"C:\Users\32891\Desktop\活动方案（第一轮）.docx")
for i, para in enumerate(doc.paragraphs):
    if para.text.strip():
        print(f"[{i}] {para.style.name}: {para.text}")
print("=== TABLES ===")
for ti, table in enumerate(doc.tables):
    print(f"\n-- Table {ti} --")
    for ri, row in enumerate(table.rows):
        cells = [cell.text.strip() for cell in row.cells]
        print(f"  Row{ri}: {cells}")
